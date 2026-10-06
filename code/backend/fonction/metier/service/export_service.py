import io
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak


def _clean(v):
    return v.strip() if isinstance(v, str) else v

def _fmt_time(v):
    return str(v)[:19].replace("T", " ") if v else ""

def _fmt_bytes(v):
    try:
        n = float(v)
    except (TypeError, ValueError):
        return ""
    for unit in ("o", "Ko", "Mo", "Go", "To"):
        if n < 1024 or unit == "To":
            return f"{n:.2f} {unit}" if unit != "o" else f"{int(n)} o"
        n /= 1024

def table_cpu(rows):
    cols = ["Date", "Utilisation (%)", "Capteur le plus chaud", "Temp. max (°C)"]
    data = []
    for r in rows:
        d = r.get("donnees") or {}
        temps = d.get("cpuTemperatureInfo") or []
        hottest = max(temps, key=lambda t: t.get("temperatureCelsius", -1)) if temps else None
        data.append([
            _fmt_time(d.get("reportTime") or r.get("report_time")),
            d.get("cpuUtilizationPct", ""),
            _clean(hottest.get("label")) if hottest else "",
            hottest.get("temperatureCelsius", "") if hottest else "",
        ])
    return cols, data

def table_ram(rows):
    cols = ["Date", "Mémoire libre", "Mémoire totale", "Page faults"]
    data = []
    for r in rows:
        d = r.get("donnees") or {}
        data.append([
            _fmt_time(d.get("reportTime") or r.get("report_time")),
            _fmt_bytes(d.get("systemRamFreeBytes")),
            _fmt_bytes(d.get("totalRamBytes")),
            d.get("pageFaults", ""),
        ])
    return cols, data

def table_stockage(rows):
    cols = ["Date", "Volume", "Total", "Libre"]
    data = []
    for r in rows:
        d = r.get("donnees") or {}
        for disk in d.get("disk") or []:
            data.append([
                _fmt_time(d.get("reportTime") or r.get("report_time")),
                disk.get("volumeId", disk.get("model", "")),
                _fmt_bytes(disk.get("storageTotalBytes", disk.get("sizeBytes"))),
                _fmt_bytes(disk.get("storageFreeBytes")),
            ])
    return cols, data

def table_batterie(rows):
    cols = ["Date", "Santé", "Cycles", "Capacité actuelle", "Capacité d'origine", "Tension"]
    data = []
    for r in rows:
        d = r.get("donnees") or {}
        data.append([
            _fmt_time(d.get("reportTime") or r.get("report_time")),
            (d.get("batteryHealth") or "").replace("BATTERY_HEALTH_", ""),
            d.get("cycleCount", ""),
            d.get("fullChargeCapacity", ""),
            d.get("designCapacity", ""),
            d.get("batteryVoltage", ""),
        ])
    return cols, data

def table_reseau(rows):
    cols = ["Date", "IP locale", "Passerelle", "Type", "État", "Débit (Kbps)"]
    data = []
    for r in rows:
        d = r.get("donnees") or {}
        data.append([
            _fmt_time(d.get("reportTime") or r.get("report_time")),
            d.get("lanIpAddress", ""),
            d.get("gatewayIpAddress", ""),
            d.get("connectionType", ""),
            d.get("connectionState", ""),
            d.get("linkDownSpeedKbps", ""),
        ])
    return cols, data

def table_peripheriques(rows):
    cols = ["Date", "Fabricant", "Nom", "VID", "PID"]
    data = []
    seen = set()
    for r in rows:
        d = r.get("donnees") or {}
        for p in d.get("usbPeripheralReport") or []:
            key = (p.get("vid"), p.get("pid"))
            if key in seen:
                continue
            seen.add(key)
            data.append([
                _fmt_time(d.get("reportTime") or r.get("report_time")),
                p.get("vendor", ""), p.get("name", ""),
                p.get("vid", ""), p.get("pid", ""),
            ])
    return cols, data


TABLE_BUILDERS = {
    "cpu": ("CPU", table_cpu),
    "ram": ("Mémoire RAM", table_ram),
    "stockage": ("Stockage", table_stockage),
    "batterie": ("Batterie", table_batterie),
    "reseau": ("Réseau", table_reseau),
    "peripheriques": ("Périphériques", table_peripheriques),
}

PROJECT_ROOT = Path(__file__).resolve().parents[5]
EXPORT_ROOT = PROJECT_ROOT / "export"
EXPORT_PDF_DIR = EXPORT_ROOT / "pdf"
EXPORT_EXCEL_DIR = EXPORT_ROOT / "excel"

for directory in (EXPORT_PDF_DIR, EXPORT_EXCEL_DIR):
    directory.mkdir(parents=True, exist_ok=True)


def sanitize_export_name(value):
    text = str(value or "appareil").strip()
    text = re.sub(r"[^a-zA-Z0-9_\- ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace(" ", "_").lower() if text else "appareil"


def resolve_device_export_name(device_id):
    try:
        from backend.fonction.conn.connexion import get_connection, close_connection
        from backend.fonction.metier.repository.devices import get_device_by_id

        connexion = get_connection()
        if connexion is None:
            return f"appareil_{device_id}"

        try:
            cur = connexion.cursor()
            device = get_device_by_id(cur, device_id)
            if not device:
                return f"appareil_{device_id}"
            return sanitize_export_name(device[3] or device[1] or f"appareil_{device_id}")
        finally:
            close_connection(connexion)
    except Exception:
        return f"appareil_{device_id}"


def build_export_filename(onglet, device_id, ext, device_name=None):
    safe_name = sanitize_export_name(device_name or resolve_device_export_name(device_id))
    return f"rapport_{safe_name}_{onglet}.{ext}"


def save_export_file(buffer, ext, filename):
    directory = EXPORT_PDF_DIR if ext == "pdf" else EXPORT_EXCEL_DIR
    file_path = directory / filename
    file_path.write_bytes(buffer.getvalue())
    return file_path


# ---------- 2. Sections : un onglet = une section, "general" = toutes ----------

def build_sections(onglet, reports):
    """
    reports : pour un onglet -> liste de rapports
              pour 'general' -> dict {cpu: [...], ram: [...], ...}
    Retourne une liste de (titre, colonnes, lignes)
    """
    if onglet == "general":
        sections = []
        for key, (title, builder) in TABLE_BUILDERS.items():
            cols, data = builder(reports.get(key) or [])
            sections.append((title, cols, data))
        return sections
    title, builder = TABLE_BUILDERS[onglet]
    cols, data = builder(reports or [])
    return [(title, cols, data)]


# ---------- 3. Générateurs ----------

def generate_excel(sections, titre_doc, device_info=None):
    wb = Workbook()
    wb.remove(wb.active)
    title_fill = PatternFill("solid", fgColor="0F172A")
    header_fill = PatternFill("solid", fgColor="0E7490")
    header_font = Font(bold=True, color="FFFFFF")
    label_font = Font(bold=True, color="334155")
    value_font = Font(bold=False, color="0F172A")

    # Onglet Fiche Matériel si device_info est fourni
    if device_info:
        ws_info = wb.create_sheet(title="Fiche Matériel")
        ws_info.sheet_view.showGridLines = True
        ws_info.merge_cells("A1:D1")
        ws_info["A1"] = "FICHE D'INFORMATION DU MATÉRIEL"
        ws_info["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        ws_info["A1"].fill = title_fill
        ws_info["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_info.row_dimensions[1].height = 28

        recents_str = ", ".join(device_info.get("utilisateurs_recents") or []) or "Aucun"

        info_data = [
            ("Catégorie", "Propriété", "Valeur", ""),
            ("Appareil", "Modèle", device_info.get("modele") or "Inconnu", ""),
            ("", "N° de Série", device_info.get("serial_number") or "N/A", ""),
            ("", "Filiale", device_info.get("filiale") or "Non attribuée", ""),
            ("", "Version OS", device_info.get("chromeos_version") or "N/A", ""),
            ("Composants", "Processeur", device_info.get("cpu_model") or "N/A", ""),
            ("", "Architecture", device_info.get("cpu_architecture") or "N/A", ""),
            ("", "Fréquence Max", device_info.get("cpu_freq_max_label") or "N/A", ""),
            ("", "RAM Totale", device_info.get("ram_total_label") or "N/A", ""),
            ("Stockage & Accès", "Modèle Disque", device_info.get("disk_model") or "N/A", ""),
            ("", "Capacité Disque", device_info.get("disk_total_label") or "N/A", ""),
            ("", "Utilisateurs Récents", recents_str, ""),
        ]

        ws_info.append([])
        for row in info_data:
            ws_info.append(list(row))

        for r in range(3, 15):
            cell_cat = ws_info.cell(row=r, column=1)
            cell_prop = ws_info.cell(row=r, column=2)
            cell_val = ws_info.cell(row=r, column=3)
            
            if r == 3:
                for c in range(1, 4):
                    cell = ws_info.cell(row=r, column=c)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell_cat.font = label_font
                cell_prop.font = label_font
                cell_val.font = value_font
                cell_val.alignment = Alignment(vertical="center")

        ws_info.column_dimensions["A"].width = 20
        ws_info.column_dimensions["B"].width = 24
        ws_info.column_dimensions["C"].width = 45

    for title, cols, data in sections:
        ws = wb.create_sheet(title=title[:31])
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(cols), 1))
        ws["A1"] = titre_doc
        ws["A1"].font = Font(bold=True, size=15, color="FFFFFF")
        ws["A1"].fill = title_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        ws.append([])
        ws.append([])
        ws.append(cols)

        for cell in ws[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style='thin', color='D1D5DB'),
                right=Side(style='thin', color='D1D5DB'),
                top=Side(style='thin', color='D1D5DB'),
                bottom=Side(style='thin', color='D1D5DB')
            )

        for row in data:
            ws.append([str(cell) if (cell is not None and str(cell).strip() != "") else "-" for cell in row])

        for i, col in enumerate(cols, start=1):
            values = [str(col)] + [str(r[i - 1]) if r[i - 1] is not None else "" for r in data]
            longest = max(len(v) for v in values) if values else 10
            ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 12), 32)

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(cols)):
            for cell in row:
                if cell.row in (1, 3):
                    continue
                cell.alignment = Alignment(vertical="center")
                cell.border = Border(
                    left=Side(style='thin', color='E5E7EB'),
                    right=Side(style='thin', color='E5E7EB'),
                    top=Side(style='thin', color='E5E7EB'),
                    bottom=Side(style='thin', color='E5E7EB')
                )

        ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{ws.max_row}"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdf(sections, titre_doc, sous_titre="", device_info=None):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=18,
        rightMargin=18,
        topMargin=18,
        bottomMargin=18,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.fontName = "Helvetica-Bold"
    title_style.fontSize = 16
    title_style.textColor = colors.HexColor("#0F172A")
    title_style.leading = 18

    subtitle_style = styles["BodyText"]
    subtitle_style.fontName = "Helvetica"
    subtitle_style.fontSize = 9
    subtitle_style.textColor = colors.HexColor("#475569")

    elements = [Paragraph(titre_doc, title_style)]
    if sous_titre:
        elements.append(Paragraph(sous_titre, subtitle_style))
    elements.append(Spacer(1, 8))

    # Bloc Fiche d'information Matériel en haut du PDF si device_info est présent
    if device_info:
        recents_str = ", ".join(device_info.get("utilisateurs_recents") or []) or "Aucun"
        
        info_table_data = [
            [Paragraph("<b>FICHE D'INFORMATION DU MATÉRIEL</b>", styles["Normal"]), "", "", ""],
            [
                Paragraph(f"<b>Modèle :</b> {device_info.get('modele') or 'Inconnu'}", styles["Normal"]),
                Paragraph(f"<b>Processeur :</b> {device_info.get('cpu_model') or 'N/A'}", styles["Normal"]),
                Paragraph(f"<b>Modèle Disque :</b> {device_info.get('disk_model') or 'N/A'}", styles["Normal"]),
                Paragraph(f"<b>RAM Totale :</b> {device_info.get('ram_total_label') or 'N/A'}", styles["Normal"]),
            ],
            [
                Paragraph(f"<b>N° de Série :</b> {device_info.get('serial_number') or 'N/A'}", styles["Normal"]),
                Paragraph(f"<b>Architecture :</b> {device_info.get('cpu_architecture') or 'N/A'}", styles["Normal"]),
                Paragraph(f"<b>Capacité Disque :</b> {device_info.get('disk_total_label') or 'N/A'}", styles["Normal"]),
                Paragraph(f"<b>Utilisateurs Récents :</b> {recents_str}", styles["Normal"]),
            ],
            [
                Paragraph(f"<b>Filiale :</b> {device_info.get('filiale') or 'Non attribuée'}", styles["Normal"]),
                Paragraph(f"<b>Fréquence Max :</b> {device_info.get('cpu_freq_max_label') or 'N/A'}", styles["Normal"]),
                Paragraph(f"<b>Version OS :</b> {device_info.get('chromeos_version') or 'N/A'}", styles["Normal"]),
                "",
            ]
        ]
        info_table = Table(info_table_data, colWidths=[190, 200, 200, 190])
        info_table.setStyle(TableStyle([
            ("SPAN", (0, 0), (3, 0)),
            ("BACKGROUND", (0, 0), (3, 0), colors.HexColor("#0F172A")),
            ("TEXTCOLOR", (0, 0), (3, 0), colors.white),
            ("ALIGN", (0, 0), (3, 0), "CENTER"),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F1F5F9")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 12))

    for _, (title, cols, data) in enumerate(sections):
        elements.append(Paragraph(title, styles["Heading2"]))
        if not data:
            elements.append(Paragraph("Aucune donnée disponible pour cette section.", styles["Italic"]))
        else:
            table_data = [[str(c) for c in cols]] + [
                [str(c) if (c is not None and str(c).strip() != "") else "-" for c in row]
                for row in data
            ]
            table = Table(table_data, repeatRows=1, colWidths=[None] * len(cols))
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E7490")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]))
            elements.append(table)
        elements.append(Spacer(1, 14))

    doc.build(elements)
    buffer.seek(0)
    return buffer